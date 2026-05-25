"""Markdown-to-DOCX converter for final review documents.

Converts the markdown output of _generate_final_review_md() into a
properly formatted .docx file using python-docx.
"""

import io
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn


# ── Regex patterns for markdown line detection ──────────────────────────

_H1_RE = re.compile(r"^# (.+)$")
_H2_RE = re.compile(r"^## (.+)$")
_H3_RE = re.compile(r"^### (.+)$")
_H4_RE = re.compile(r"^#### (.+)$")
_BOLD_LABEL_RE = re.compile(r"^\*\*(.+?)\*\*:\s*(.*)$")
_QUOTE_RE = re.compile(r'^"(.*)"$')


# ── Font name per language ──────────────────────────────────────────────

def _font_name(lang: str) -> str:
    """Return the preferred font name for the given language."""
    if lang == "ja":
        return "Yu Mincho"
    return "Times New Roman"


# ── Low-level font helpers ──────────────────────────────────────────────

def _set_run_font(run, lang: str, size: Pt, bold: bool = False,
                  italic: bool = False):
    """Configure font properties on a Run object, including East-Asian font."""
    font = run.font
    font.name = _font_name(lang)
    font.size = size
    font.bold = bold
    font.italic = italic
    # Set East-Asian font face via XML (python-docx high-level API lacks this)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = run._element.makeelement(qn("w:rFonts"), {})
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:eastAsia"), _font_name(lang))


def _set_style_font(style, lang: str, size: Pt):
    """Configure Normal style font for the document."""
    font = style.font
    font.name = _font_name(lang)
    font.size = size
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = style.element.makeelement(qn("w:rFonts"), {})
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:eastAsia"), _font_name(lang))


# ── Page and document setup ─────────────────────────────────────────────

def _setup_page(doc):
    """Configure A4 page with 25.4mm margins."""
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)


def _define_styles(doc, lang: str):
    """Set up document-level default styles."""
    style = doc.styles["Normal"]
    _set_style_font(style, lang, Pt(10.5))
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.line_spacing = 1.15


# ── Paragraph builders ──────────────────────────────────────────────────

def _add_heading_para(doc, text: str, level: int, lang: str):
    """Add a heading paragraph. Sizes: H1=16pt, H2=14pt, H3=12pt, H4=11pt."""
    sizes = {1: Pt(16), 2: Pt(14), 3: Pt(12), 4: Pt(11)}
    size = sizes.get(level, Pt(10.5))
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(12)
    para.paragraph_format.space_after = Pt(4)
    run = para.add_run(text)
    _set_run_font(run, lang, size, bold=True)
    return para


def _add_body_para(doc, text: str, lang: str):
    """Add a normal body text paragraph (10.5pt)."""
    para = doc.add_paragraph()
    run = para.add_run(text)
    _set_run_font(run, lang, Pt(10.5))
    return para


def _add_quote_para(doc, text: str, lang: str):
    """Add a quoted excerpt paragraph (italic, indented 1cm)."""
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Cm(1.0)
    para.paragraph_format.right_indent = Cm(1.0)
    run = para.add_run(text)
    _set_run_font(run, lang, Pt(10), italic=True)
    return para


def _add_labeled_para(doc, label: str, body: str, lang: str):
    """Add a bold-label + normal-text paragraph.

    Example: **Recommendation**: Major Revision
    -> [Bold "Recommendation:"] [Normal " Major Revision"]
    """
    para = doc.add_paragraph()
    run_label = para.add_run(label + ": ")
    _set_run_font(run_label, lang, Pt(10.5), bold=True)
    if body:
        run_body = para.add_run(body)
        _set_run_font(run_body, lang, Pt(10.5))
    return para


# ── Main converter ──────────────────────────────────────────────────────

def convert_md_to_docx(md_text: str, lang: str) -> bytes:
    """Convert a final review markdown string to a DOCX byte stream.

    Parses the markdown line by line with a simple state machine and
    builds a python-docx Document object.

    Args:
        md_text: The markdown content from _generate_final_review_md().
        lang: Language code ("en" or "ja").

    Returns:
        bytes: The serialized .docx file content.
    """
    doc = Document()
    _setup_page(doc)
    _define_styles(doc, lang)

    lines = md_text.split("\n")
    i = 0
    n = len(lines)

    while i < n:
        line = lines[i]
        stripped = line.strip()

        # ── Blank line: skip ──
        if not stripped:
            i += 1
            continue

        # ── H1 (# heading) ──
        m = _H1_RE.match(stripped)
        if m:
            _add_heading_para(doc, m.group(1), 1, lang)
            i += 1
            continue

        # ── H2 (## heading) ──
        m = _H2_RE.match(stripped)
        if m:
            _add_heading_para(doc, m.group(1), 2, lang)
            i += 1
            continue

        # ── H3 (### heading) ──
        m = _H3_RE.match(stripped)
        if m:
            _add_heading_para(doc, m.group(1), 3, lang)
            i += 1
            continue

        # ── H4 (#### heading) ──
        m = _H4_RE.match(stripped)
        if m:
            _add_heading_para(doc, m.group(1), 4, lang)
            i += 1
            continue

        # ── Quoted excerpt ("...") ──
        m = _QUOTE_RE.match(stripped)
        if m:
            _add_quote_para(doc, m.group(1), lang)
            i += 1
            continue

        # ── Bold label + text (**Label**: text) ──
        m = _BOLD_LABEL_RE.match(stripped)
        if m:
            _add_labeled_para(doc, m.group(1), m.group(2), lang)
            i += 1
            continue

        # ── Regular paragraph: accumulate until blank line ──
        para_lines = []
        while i < n and lines[i].strip():
            para_lines.append(lines[i].strip())
            i += 1
        if para_lines:
            _add_body_para(doc, " ".join(para_lines), lang)

    # Serialize to bytes
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
